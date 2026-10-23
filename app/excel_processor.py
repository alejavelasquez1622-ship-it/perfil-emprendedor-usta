import hashlib
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional
import openpyxl
from sqlalchemy.orm import Session
from app.config import DATA_DIR, PROFILE_QUESTIONS, INNPULSA_STAGES
from app.models import SurveyResponse, AcademicPeriod, Faculty, LoadHistory

# Canonical faculty names for standardization
CANONICAL_FACULTIES = {
    "ingenieria de sistemas": "Ingeniería de Sistemas",
    "ingeniería de sistemas": "Ingeniería de Sistemas",
    "sistemas": "Ingeniería de Sistemas",
    "contaduria publica": "Contaduría Pública",
    "contaduría pública": "Contaduría Pública",
    "contaduria": "Contaduría Pública",
    "ingenieria mecatronica": "Ingeniería Mecatrónica",
    "ingeniería mecatrónica": "Ingeniería Mecatrónica",
    "mecatronica": "Ingeniería Mecatrónica",
    "derecho": "Derecho",
    "arquitectura": "Arquitectura",
    "ingenieria industrial": "Ingeniería Industrial",
    "ingeniería industrial": "Ingeniería Industrial",
    "industrial": "Ingeniería Industrial",
    "ingenieria ambiental": "Ingeniería Ambiental",
    "ingeniería ambiental": "Ingeniería Ambiental",
    "ambiental": "Ingeniería Ambiental",
    "negocios internacionales": "Negocios Internacionales",
    "negocios": "Negocios Internacionales",
    "arte y cultura": "Arte y Cultura",
    "psicologia": "Psicología",
    "psicología": "Psicología",
    "administracion de empresas agropecuarias": "Administración de Empresas Agropecuarias",
    "administración de empresas agropecuarias": "Administración de Empresas Agropecuarias",
    "agropecuarias": "Administración de Empresas Agropecuarias",
    "cultura fisica, deporte y recreacion": "Cultura Física, Deporte y Recreación",
    "cultura física, deporte y recreación": "Cultura Física, Deporte y Recreación",
    "cultura fisica": "Cultura Física, Deporte y Recreación",
    "ingenieria civil": "Ingeniería Civil",
    "ingeniería civil": "Ingeniería Civil",
    "civil": "Ingeniería Civil",
    "lenguas extranjeras": "Lenguas Extranjeras",
    "lenguas": "Lenguas Extranjeras"
}

def clean_text(val: Any) -> str:
    """Normalize whitespace and strip string."""
    if val is None:
        return ""
    text = str(val).replace("\xa0", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text

def normalize_faculty(fac: str) -> str:
    clean = clean_text(fac)
    if not clean:
        return "Sin Facultad Especificada"
    low = clean.lower()
    for key, canonical in CANONICAL_FACULTIES.items():
        if key in low:
            return canonical
    return clean

def normalize_stage(stage_val: str) -> str:
    clean = clean_text(stage_val)
    if not clean:
        return "Sin etapa identificada"
    for stage in INNPULSA_STAGES:
        if clean.lower().startswith(stage.lower()) or stage.lower() in clean.lower():
            return stage
    return "Sin etapa identificada"

def detect_period_from_date(date_str: str) -> str:
    """Detects period e.g. '2026-04' -> '2026-1', '2026-09' -> '2026-2'."""
    clean = clean_text(date_str)
    if not clean:
        return "2026-1"
    match = re.search(r"(\d{4})[-/](\d{1,2})", clean)
    if match:
        year = int(match.group(1))
        month = int(match.group(2))
        sem = 1 if month <= 6 else 2
        return f"{year}-{sem}"
    # Fallback search for 4-digit year
    match_year = re.search(r"\b(20\d{2})\b", clean)
    if match_year:
        return f"{match_year.group(1)}-1"
    return "2026-1"

def calculate_row_hash(row_data: Dict[str, Any]) -> str:
    """Creates deterministic SHA-256 hash of row contents for change detection."""
    sorted_items = sorted([(k, str(v)) for k, v in row_data.items() if k != "raw_data_hash"])
    raw_str = "|".join(f"{k}:{v}" for k, v in sorted_items)
    return hashlib.sha256(raw_str.encode("utf-8")).hexdigest()

class ExcelProcessor:
    TEMP_DIR = DATA_DIR / "temp_uploads"

    def __init__(self):
        self.TEMP_DIR.mkdir(parents=True, exist_ok=True)

    def parse_excel_file(self, file_path: Path) -> List[Dict[str, Any]]:
        """Parses Microsoft Forms .xlsx or .csv into standardized dictionaries."""
        if str(file_path).lower().endswith(".csv"):
            return self._parse_csv_file(file_path)
        return self._parse_xlsx_file(file_path)

    def _parse_csv_file(self, file_path: Path) -> List[Dict[str, Any]]:
        import csv
        for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    sample = f.read(4096)
                    f.seek(0)
                    dialect = csv.Sniffer().sniff(sample, delimiters=";,")
                    reader = csv.reader(f, dialect)
                    rows_raw = list(reader)
                    if rows_raw and len(rows_raw) > 1:
                        return self._process_raw_rows(rows_raw[0], rows_raw[1:])
            except Exception:
                continue

        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            rows_raw = list(reader)
            if rows_raw and len(rows_raw) > 1:
                return self._process_raw_rows(rows_raw[0], rows_raw[1:])
        return []

    def _parse_xlsx_file(self, file_path: Path) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        if "Sheet1" in wb.sheetnames:
            ws = wb["Sheet1"]

        headers_raw = [cell.value for cell in ws[1]]
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        return self._process_raw_rows(headers_raw, data_rows)

    def _process_raw_rows(self, headers_raw: List[Any], data_rows: List[Any]) -> List[Dict[str, Any]]:
        headers_clean = [clean_text(h) for h in headers_raw]

        rows = []
        for row in data_rows:
            if not any(row):
                continue
            row_dict = {}
            funding_seen = 0
            for idx, col_val in enumerate(row):
                if idx >= len(headers_clean):
                    continue
                header = headers_clean[idx]
                h_low = header.lower().strip()
                val = clean_text(col_val)

                # Map standard columns
                if h_low in ["id", "id.", "id de respuesta", "identificador"]:
                    row_dict["original_id"] = val
                elif "hora de inicio" in h_low:
                    row_dict["submission_start"] = val
                elif "hora de finalización" in h_low or "hora de finalizacion" in h_low:
                    row_dict["submission_end"] = val
                elif "autorización expresa" in h_low or "autorizacion expresa" in h_low or "tratamiento de datos" in h_low:
                    row_dict["authorized"] = "autorizo" in val.lower() and "no autorizo" not in val.lower()
                    row_dict["auth_raw"] = val
                elif "facultad" in h_low:
                    row_dict["faculty_name"] = normalize_faculty(val)
                elif "semestre" in h_low:
                    row_dict["semester_enrolled"] = val
                # Profile Questions
                elif "punto de vista de la motivación" in h_low or "punto de vista de la motivacion" in h_low:
                    row_dict["q_motivacion"] = val
                elif "problemas que plantea la gestión" in h_low or "problemas que plantea la gestion" in h_low:
                    row_dict["q_gestion_problemas"] = val
                elif "motivación económica es importante" in h_low or "motivacion economica es importante" in h_low:
                    row_dict["q_motivacion_economica"] = val
                elif "disponer de dinero propio" in h_low:
                    row_dict["q_dinero_propio"] = val
                elif "siguientes definiciones se siente mas identificado" in h_low or "siguientes definiciones se siente más identificado" in h_low:
                    row_dict["q_iniciativa"] = val  # P13
                elif "de cara al nuevo proyecto" in h_low:
                    row_dict["q_nuevo_proyecto"] = val
                elif "tiempo que debe dedicar" in h_low:
                    row_dict["q_tiempo_dedicacion"] = val
                elif "relación e influencia social" in h_low or "relacion e influencia social" in h_low:
                    row_dict["q_influencia_social"] = val  # P16
                elif "función comercial" in h_low or "funcion comercial" in h_low:
                    row_dict["q_funcion_comercial"] = val  # P17
                elif "dotes de liderazgo" in h_low:
                    row_dict["q_liderazgo"] = val
                elif "haber acertado con una solución" in h_low or "haber acertado con una solucion" in h_low:
                    row_dict["q_satisfaccion_logros"] = val
                elif "arriesgarse al tomar decisiones" in h_low:
                    row_dict["q_actitud_riesgo"] = val  # P20
                elif "idea que le va a permitir montar una empresa" in h_low:
                    row_dict["q_validacion_viabilidad"] = val  # P21
                # Initiative Status
                elif "posee una idea de negocio o emprendimiento" in h_low:
                    row_dict["initiative_status"] = val
                    row_dict["has_initiative"] = "sí" in val.lower() or "si" in val.lower()
                elif "etapa del emprendimiento según innpulsa" in h_low or "etapa del emprendimiento segun innpulsa" in h_low:
                    row_dict["stage"] = normalize_stage(val)
                    row_dict["stage_raw"] = val
                elif "nombre de la idea emprendedora" in h_low:
                    row_dict["idea_name"] = val
                elif "descripción breve de la idea" in h_low or "descripcion breve de la idea" in h_low:
                    row_dict["idea_description"] = val
                elif "problema o necesidad específica" in h_low or "problema o necesidad especifica" in h_low:
                    row_dict["problem_addressed"] = val
                elif "tiempo ha invertido en el proyecto" in h_low:
                    row_dict["time_invested_months"] = val
                elif "ha invertido recursos económicos" in h_low or "ha invertido recursos economicos" in h_low:
                    row_dict["has_invested_money"] = val
                elif "monto invertido (cifras)" in h_low or "monto invertido" in h_low:
                    row_dict["invested_amount_raw"] = val
                elif "cómo obtuviste el recurso económico" in h_low or "como obtuviste el recurso economico" in h_low:
                    funding_seen += 1
                    if funding_seen == 1:
                        row_dict["funding_source_1"] = val
                    else:
                        row_dict["funding_source_2"] = val
                elif "ventas mensuales promedio" in h_low:
                    row_dict["monthly_sales_raw"] = val
                elif "desarrolla la idea de negocio" in h_low:
                    row_dict["business_model_type"] = val
                elif "integrantes conforman el proyecto" in h_low:
                    row_dict["team_size_raw"] = val
                elif "recibido apoyo de instituciones" in h_low:
                    row_dict["received_support"] = val
                elif "instituciones de las cuales ha recibido apoyo" in h_low:
                    row_dict["support_entities_raw"] = val
                elif "sector pertenece su emprendimiento" in h_low:
                    row_dict["economic_sector"] = val
                elif "mercado objetivo" in h_low:
                    row_dict["target_market"] = val

            # Fill defaults if missing
            if "original_id" not in row_dict or not row_dict["original_id"]:
                row_dict["original_id"] = str(len(rows) + 1)
            if "authorized" not in row_dict:
                row_dict["authorized"] = True
            if "has_initiative" not in row_dict:
                row_dict["has_initiative"] = False
            if "faculty_name" not in row_dict or not row_dict["faculty_name"]:
                row_dict["faculty_name"] = "Sin Facultad"

            # Compute hash
            row_dict["raw_data_hash"] = calculate_row_hash(row_dict)
            rows.append(row_dict)

        return rows

    def analyze_preview(
        self,
        rows: List[Dict[str, Any]],
        db: Session,
        filename: str,
        target_period_code: Optional[str] = None
    ) -> Dict[str, Any]:
        """Analyzes rows against the database and returns preview metrics."""
        warnings = []

        # Detect period from date if not specified
        periods_detected = {}
        for r in rows:
            p = detect_period_from_date(r.get("submission_start", ""))
            periods_detected[p] = periods_detected.get(p, 0) + 1

        dominant_period = target_period_code or max(periods_detected.items(), key=lambda x: x[1])[0]

        if len(periods_detected) > 1 and not target_period_code:
            warnings.append(
                f"El archivo contiene fechas correspondientes a múltiples periodos: {dict(periods_detected)}. "
                f"Se sugiere asociar al periodo dominante '{dominant_period}', o confirmar el periodo deseado."
            )

        # Check existing period in DB
        if target_period_code and target_period_code not in ["auto_split", "Todos", "auto"]:
            period_obj = db.query(AcademicPeriod).filter(AcademicPeriod.code == target_period_code).first()
            existing_responses = db.query(SurveyResponse).filter(SurveyResponse.period_id == period_obj.id).all() if period_obj else []
        else:
            detected_codes = list(periods_detected.keys())
            existing_responses = db.query(SurveyResponse).filter(SurveyResponse.period_code.in_(detected_codes)).all()

        current_db_count = len(existing_responses)
        existing_map = {}
        for ex in existing_responses:
            existing_map[str(ex.original_id)] = ex.raw_data_hash

        total_rows = len(rows)
        authorized_rows = sum(1 for r in rows if r.get("authorized", False))
        unauthorized_rows = total_rows - authorized_rows

        new_count = 0
        updated_count = 0
        unchanged_count = 0

        for r in rows:
            orig_id = str(r["original_id"])
            if orig_id not in existing_map:
                new_count += 1
            else:
                if existing_map[orig_id] == r["raw_data_hash"]:
                    unchanged_count += 1
                else:
                    updated_count += 1

        is_fewer = current_db_count > 0 and total_rows < current_db_count
        if is_fewer:
            warnings.append(
                f"¡Advertencia institucional! El archivo cargado contiene {total_rows} registros, "
                f"lo cual es MENOR a los {current_db_count} registros ya registrados en el periodo {dominant_period}. "
                "Verifique si se trata de un corte parcial o una sustitución."
            )

        faculties = sorted(list(set(r.get("faculty_name", "Sin Facultad") for r in rows if r.get("authorized", False))))

        # Save parsed data to temp cache file
        temp_id = hashlib.md5(f"{filename}_{datetime.utcnow().isoformat()}".encode("utf-8")).hexdigest()
        temp_path = self.TEMP_DIR / f"{temp_id}.json"
        with open(temp_path, "w", encoding="utf-8") as f:
            json.dump({
                "filename": filename,
                "detected_period": dominant_period,
                "rows": rows
            }, f, ensure_ascii=False)

        preview_samples = []
        for r in rows[:8]:
            preview_samples.append({
                "original_id": r["original_id"],
                "faculty": r.get("faculty_name", ""),
                "semester": r.get("semester_enrolled", ""),
                "authorized": "Autorizado" if r.get("authorized") else "No Autorizado",
                "initiative": r.get("initiative_status", "Sin respuesta"),
                "stage": r.get("stage", "N/A"),
                "status": "NUEVO" if str(r["original_id"]) not in existing_map else (
                    "SIN CAMBIOS" if existing_map[str(r["original_id"])] == r["raw_data_hash"] else "MODIFICADO"
                )
            })

        return {
            "filename": filename,
            "detected_period": dominant_period,
            "periods_detected": periods_detected,
            "total_rows": total_rows,
            "authorized_rows": authorized_rows,
            "unauthorized_rows": unauthorized_rows,
            "new_rows": new_count,
            "updated_rows": updated_count,
            "unchanged_rows": unchanged_count,
            "current_db_count": current_db_count,
            "is_fewer_than_current": is_fewer,
            "faculties_found": faculties,
            "warnings": warnings,
            "preview_samples": preview_samples,
            "temp_file_id": temp_id
        }

    def commit_upload(
        self,
        temp_file_id: str,
        period_code: str,
        db: Session,
        loaded_by: str = "Administrador",
        mode: str = "smart_update"
    ) -> Dict[str, Any]:
        """Commits the verified upload to the database safely."""
        temp_path = self.TEMP_DIR / f"{temp_file_id}.json"
        if not temp_path.exists():
            raise ValueError("El archivo temporal de vista previa ha expirado o no existe. Cargue el archivo nuevamente.")

        with open(temp_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        rows = data["rows"]
        filename = data["filename"]

        def get_or_create_period(p_code: str) -> AcademicPeriod:
            p_obj = db.query(AcademicPeriod).filter(AcademicPeriod.code == p_code).first()
            if not p_obj:
                match = re.match(r"^(\d{4})-(\d)$", p_code)
                y = int(match.group(1)) if match else 2026
                s = int(match.group(2)) if match else 1
                p_obj = AcademicPeriod(
                    code=p_code,
                    year=y,
                    semester=s,
                    description=f"Periodo Académico {p_code}",
                    is_active=True
                )
                db.add(p_obj)
                db.commit()
                db.refresh(p_obj)
            return p_obj

        use_auto_split = (period_code in ["auto_split", "auto", "detect", "auto_distribuir"])

        # Track existing records and faculties
        existing_faculties = {f.name for f in db.query(Faculty).all()}
        periods_cache = {}

        if mode == "replace_all":
            if not use_auto_split:
                p_obj = get_or_create_period(period_code)
                periods_cache[period_code] = p_obj
                db.query(SurveyResponse).filter(SurveyResponse.period_id == p_obj.id).delete()
            else:
                detected_period_codes = set(detect_period_from_date(r.get("submission_start", "")) for r in rows)
                for p_c in detected_period_codes:
                    p_o = db.query(AcademicPeriod).filter(AcademicPeriod.code == p_c).first()
                    if p_o:
                        db.query(SurveyResponse).filter(SurveyResponse.period_id == p_o.id).delete()
            db.commit()
        elif not use_auto_split:
            p_obj = get_or_create_period(period_code)
            periods_cache[period_code] = p_obj

        new_count = 0
        updated_count = 0
        unchanged_count = 0

        for r in rows:
            orig_id = str(r["original_id"])
            fac_name = r.get("faculty_name", "Sin Facultad")

            # Determine period for this row
            if use_auto_split:
                row_pcode = detect_period_from_date(r.get("submission_start", ""))
            else:
                row_pcode = period_code

            if row_pcode not in periods_cache:
                periods_cache[row_pcode] = get_or_create_period(row_pcode)
            target_period_obj = periods_cache[row_pcode]

            # Register faculty if new
            if fac_name and fac_name not in existing_faculties:
                new_fac = Faculty(name=fac_name)
                db.add(new_fac)
                existing_faculties.add(fac_name)

            if use_auto_split:
                existing_resp = db.query(SurveyResponse).filter(
                    SurveyResponse.original_id == orig_id
                ).first()
            else:
                existing_resp = db.query(SurveyResponse).filter(
                    SurveyResponse.period_id == target_period_obj.id,
                    SurveyResponse.original_id == orig_id
                ).first()

            if existing_resp:
                if existing_resp.raw_data_hash == r["raw_data_hash"]:
                    unchanged_count += 1
                else:
                    for key, val in r.items():
                        if hasattr(existing_resp, key) and key not in ["id", "period_id"]:
                            setattr(existing_resp, key, val)
                    existing_resp.period_id = target_period_obj.id
                    existing_resp.period_code = row_pcode
                    existing_resp.updated_at = datetime.utcnow()
                    updated_count += 1
            else:
                new_resp = SurveyResponse(
                    period_id=target_period_obj.id,
                    period_code=row_pcode,
                    original_id=orig_id,
                    submission_start=r.get("submission_start"),
                    submission_end=r.get("submission_end"),
                    authorized=r.get("authorized", True),
                    faculty_name=fac_name,
                    program_name=r.get("program_name"),
                    semester_enrolled=r.get("semester_enrolled"),
                    q_motivacion=r.get("q_motivacion"),
                    q_gestion_problemas=r.get("q_gestion_problemas"),
                    q_motivacion_economica=r.get("q_motivacion_economica"),
                    q_dinero_propio=r.get("q_dinero_propio"),
                    q_iniciativa=r.get("q_iniciativa"),
                    q_nuevo_proyecto=r.get("q_nuevo_proyecto"),
                    q_tiempo_dedicacion=r.get("q_tiempo_dedicacion"),
                    q_influencia_social=r.get("q_influencia_social"),
                    q_funcion_comercial=r.get("q_funcion_comercial"),
                    q_liderazgo=r.get("q_liderazgo"),
                    q_satisfaccion_logros=r.get("q_satisfaccion_logros"),
                    q_actitud_riesgo=r.get("q_actitud_riesgo"),
                    q_validacion_viabilidad=r.get("q_validacion_viabilidad"),
                    has_initiative=r.get("has_initiative", False),
                    initiative_status=r.get("initiative_status"),
                    stage=r.get("stage"),
                    stage_raw=r.get("stage_raw"),
                    idea_name=r.get("idea_name"),
                    idea_description=r.get("idea_description"),
                    problem_addressed=r.get("problem_addressed"),
                    time_invested_months=r.get("time_invested_months"),
                    has_invested_money=r.get("has_invested_money"),
                    invested_amount_raw=r.get("invested_amount_raw"),
                    funding_source_1=r.get("funding_source_1"),
                    funding_source_2=r.get("funding_source_2"),
                    monthly_sales_raw=r.get("monthly_sales_raw"),
                    business_model_type=r.get("business_model_type"),
                    team_size_raw=r.get("team_size_raw"),
                    received_support=r.get("received_support"),
                    support_entities_raw=r.get("support_entities_raw"),
                    economic_sector=r.get("economic_sector"),
                    target_market=r.get("target_market"),
                    raw_data_hash=r.get("raw_data_hash")
                )
                db.add(new_resp)
                new_count += 1

        # Determine main period object for history tracking
        primary_period_obj = periods_cache[period_code] if (not use_auto_split and period_code in periods_cache) else (list(periods_cache.values())[0] if periods_cache else get_or_create_period("2026-1"))
        display_code = ", ".join(sorted(periods_cache.keys())) if use_auto_split else period_code

        # Record load history
        history = LoadHistory(
            period_id=primary_period_obj.id,
            period_code=display_code,
            filename=filename,
            total_rows=len(rows),
            authorized_rows=sum(1 for r in rows if r.get("authorized", False)),
            unauthorized_rows=sum(1 for r in rows if not r.get("authorized", False)),
            new_rows=new_count,
            updated_rows=updated_count,
            unchanged_rows=unchanged_count,
            loaded_by=loaded_by,
            loaded_at=datetime.utcnow()
        )
        db.add(history)
        for p_item in periods_cache.values():
            p_item.updated_at = datetime.utcnow()
        db.commit()

        # Remove temp file
        if temp_path.exists():
            temp_path.unlink()

        return {
            "period_code": display_code,
            "total_processed": len(rows),
            "new_records": new_count,
            "updated_records": updated_count,
            "unchanged_records": unchanged_count,
            "message": f"Carga completada exitosamente ({display_code})."
        }
